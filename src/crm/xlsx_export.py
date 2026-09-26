"""
XLSX-выгрузка ответивших продавцов: карточка + последняя переписка.

openpyxl обязателен (есть в requirements.txt). Файл собирается в памяти и
отдаётся дашбордом как attachment.
"""

from io import BytesIO
from typing import List

HEADERS = [
    ("seller_id", "seller_id"),
    ("shop_name", "Магазин"),
    ("company_name", "Юрлицо"),
    ("province", "Провинция"),
    ("orders", "Заказы"),
    ("reviews", "Отзывы"),
    ("rating", "Рейтинг"),
    ("score", "Оценка"),
    ("status", "Статус"),
    ("reply_category", "Категория ответа"),
    ("reply_at", "Ответ получен"),
    ("handed_to_manager_at", "Передан менеджеру"),
    ("stage", "Стадия"),
    ("notes", "Заметка"),
    ("last_incoming", "Последнее входящее"),
    ("last_outgoing", "Наше последнее"),
    ("seller_profile_url", "Профиль"),
]

STAGE_TITLES = {"negotiation": "переговоры", "deal": "сделка", "dead": "закрыли"}
STATUS_TITLES = {
    "new": "не написал", "sending": "отправляется", "sent": "написал",
    "replied": "ответил", "refused": "отказ", "failed": "ошибка",
    "skipped": "пропущен", "opted_out": "не писать",
}
REPLY_TITLES = {
    "positive": "ПОЗИТИВНЫЙ", "negative": "НЕГАТИВНЫЙ", "neutral": "НЕЙТРАЛЬНЫЙ",
    "interested": "ЗАИНТЕРЕСОВАН", "needs_info": "НУЖНА ДОПОЛНИТЕЛЬНАЯ ИНФОРМАЦИЯ",
    "other": "ДРУГОЕ",
}


def replied_xlsx(rows: List[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    sheet = book.active
    sheet.title = "Ответившие"

    header_fill = PatternFill("solid", fgColor="1F2937")
    for col, (_, title) in enumerate(HEADERS, 1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    sheet.freeze_panes = "A2"

    for index, row in enumerate(rows, 2):
        for col, (key, _) in enumerate(HEADERS, 1):
            value = row.get(key) or ""
            if key == "stage":
                value = STAGE_TITLES.get(value, value)
            elif key == "status":
                value = STATUS_TITLES.get(value, value)
            elif key == "reply_category":
                value = REPLY_TITLES.get(value, value)
            sheet.cell(row=index, column=col, value=value)

    widths = {"A": 12, "B": 26, "C": 30, "D": 12, "E": 9, "F": 9, "G": 8,
              "H": 7, "I": 12, "J": 30, "K": 17, "L": 17, "M": 12, "N": 28,
              "O": 46, "P": 46, "Q": 34}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
